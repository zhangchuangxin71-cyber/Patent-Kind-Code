#!/usr/bin/env python3
"""Build the Chinese patent-experiment collection Word document from final sources."""
from __future__ import annotations

import argparse
import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


MARKDOWN_FILES = (
    "专利撰写AI实验交接稿.md",
    "专利实验结果与过程说明.md",
    "新增专利对比实验详细报告.md",
    "大模型语义信息生成与实验参与详细说明.md",
    "数据划分、文本时间边界与扩散接入LightGCN技术说明.md",
)
JSON_FILE = "专利实验机器可读摘要.json"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def clean_inline(text: str) -> str:
    return re.sub(r"`([^`]*)`", r"\1", text).replace("**", "").replace("__", "")


def add_inline(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run(clean_inline(part))
        if part.startswith("**") and part.endswith("**"):
            run.bold = True
        if part.startswith("`") and part.endswith("`"):
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")


def split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def add_markdown_table(document: Document, lines: list[str]) -> None:
    rows = [split_table_row(line) for line in lines if not is_table_separator(line)]
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for col, value in enumerate(row):
            cells[col].text = clean_inline(value)
            for paragraph in cells[col].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
            if idx == 0:
                set_cell_shading(cells[col], "D9EAF7")
                for run in cells[col].paragraphs[0].runs:
                    run.bold = True
    document.add_paragraph()


def add_markdown(document: Document, content: str) -> None:
    lines = content.splitlines()
    index = 0
    in_code = False
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            in_code = not in_code
            index += 1
            continue
        if in_code:
            paragraph = document.add_paragraph(style="No Spacing")
            run = paragraph.add_run(line)
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
            run.font.size = Pt(8)
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and is_table_separator(lines[index + 1]):
            end = index + 2
            while end < len(lines) and lines[end].startswith("|"):
                end += 1
            add_markdown_table(document, lines[index:end])
            index = end
            continue
        if not line.strip():
            index += 1
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            level = min(len(heading.group(1)), 4)
            document.add_heading(clean_inline(heading.group(2)), level=level)
        elif line.startswith("> "):
            paragraph = document.add_paragraph(style="Quote")
            add_inline(paragraph, line[2:])
        elif re.match(r"^[-*+]\s+", line):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, re.sub(r"^[-*+]\s+", "", line))
        elif re.match(r"^\d+\.\s+", line):
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\.\s+", "", line))
        else:
            paragraph = document.add_paragraph()
            add_inline(paragraph, line)
        index += 1


def apply_document_style(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    source_dir = Path(args.source_dir)
    out = Path(args.out)
    document = Document()
    apply_document_style(document)

    document.add_heading("专利实验报告合集", level=0)
    document.add_paragraph("用于专利技术交底书与专利撰写 AI 的统一实验材料")
    document.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    document.add_heading("本次修订说明", level=1)
    for note in (
        "Yelp2018 与 Food 扩展表的 12 个标准差已按五个 OOD seed 的样本标准差（ddof=1）更正；均值、配对差和方向不变。",
        "corrected-v5 三方法审计现为 16 项可直接核验条件。Yelp2018、Food、Amazon Beauty 的 alpha=0.75 是固定组合参数，未提供外部语义权重确认文件，不表述为独立 validation confirmation。",
        "三方法表只支持组合方法在本次五个 seed、相应数据划分与 BPR25 预算下的整体方向；不拆分归因给因果干预、InfoNCE、语义条件或最终分数融合。",
    ):
        document.add_paragraph(note, style="List Bullet")

    for filename in MARKDOWN_FILES:
        path = source_dir / filename
        if not path.exists():
            raise FileNotFoundError(path)
        document.add_page_break()
        document.add_heading(f"汇编文件：{filename}", level=0)
        add_markdown(document, path.read_text(encoding="utf-8"))

    json_path = source_dir / JSON_FILE
    if not json_path.exists():
        raise FileNotFoundError(json_path)
    parsed = json.loads(json_path.read_text(encoding="utf-8"))
    document.add_page_break()
    document.add_heading(f"汇编文件：{JSON_FILE}", level=0)
    document.add_paragraph("以下为机器可读摘要的完整 JSON 内容，保留以便另一套 AI 或审稿人复核。")
    formatted = json.dumps(parsed, ensure_ascii=False, indent=2)
    for line in formatted.splitlines():
        paragraph = document.add_paragraph(style="No Spacing")
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
        run.font.size = Pt(7)

    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(out)
    print(out)


if __name__ == "__main__":
    main()
